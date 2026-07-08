"""Tests for the export service and the CSV / JSON / Excel exporters."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from herpetoid.application.export import ExportData, ExportService, observation_rows
from herpetoid.domain import Individual, Observation, Project, Species
from herpetoid.infrastructure.exporters import (
    CsvExporter,
    ExcelExporter,
    JsonExporter,
    default_exporters,
)


def _data() -> ExportData:
    return ExportData(
        project=Project(name="Study", uuid="u1", id=1),
        species=[Species(scientific_name="Calotriton asper", id=1)],
        individuals=[Individual(species_id=1, code="CA-001", id=1)],
        observations=[
            Observation(
                species_id=1,
                individual_id=1,
                observer="AL",
                measurements={"svl": 50.0, "sex": "female"},
                id=1,
            ),
            Observation(
                species_id=1,
                individual_id=None,
                observer="BM",
                measurements={"svl": 30.0},
                id=2,
            ),
        ],
    )


def test_observation_rows() -> None:
    columns, rows = observation_rows(_data())
    assert columns[0] == "observation_id"
    assert "svl" in columns and "sex" in columns
    assert len(rows) == 2
    assert rows[0]["individual"] == "CA-001"
    assert rows[0]["species"] == "Calotriton asper"
    assert rows[0]["svl"] == 50.0
    assert rows[1]["individual"] == ""  # unidentified observation


def test_csv_export(tmp_path: Path) -> None:
    destination = tmp_path / "out.csv"
    CsvExporter().export(_data(), destination)
    with destination.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2
    assert rows[0]["individual"] == "CA-001"
    assert rows[0]["svl"] == "50.0"


def test_json_export(tmp_path: Path) -> None:
    destination = tmp_path / "out.json"
    JsonExporter().export(_data(), destination)
    payload = json.loads(destination.read_text(encoding="utf-8"))
    assert payload["project"]["name"] == "Study"
    assert len(payload["individuals"]) == 1
    assert len(payload["observations"]) == 2


def test_excel_export(tmp_path: Path) -> None:
    destination = tmp_path / "out.xlsx"
    ExcelExporter().export(_data(), destination)
    workbook = load_workbook(destination)
    assert set(workbook.sheetnames) == {"Observations", "Individuals"}
    header = [cell.value for cell in workbook["Observations"][1]]
    assert "svl" in header
    assert workbook["Observations"].max_row == 3  # header + 2 observations


def test_export_service_dispatch(tmp_path: Path) -> None:
    service = ExportService(default_exporters())
    assert service.available_formats() == ["csv", "json", "xlsx"]
    service.export("csv", _data(), tmp_path / "d.csv")
    assert (tmp_path / "d.csv").exists()
    with pytest.raises(KeyError):
        service.export("pdf", _data(), tmp_path / "x.pdf")
