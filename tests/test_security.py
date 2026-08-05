"""Security regression tests.

HerpetoID's two untrusted inputs are (a) **project bundles**, which are portable by design and get
mailed between researchers, and (b) the **image files** a bundle imports. These tests pin the
defences against a hostile bundle or image, and against exports that a spreadsheet would execute.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from PIL import Image as PilImage

from herpetoid.application.export import ExportData
from herpetoid.application.project_service import ProjectService
from herpetoid.domain import Individual, Observation, Project, Species
from herpetoid.infrastructure.exporters import CsvExporter, ExcelExporter, JsonExporter
from herpetoid.infrastructure.image_store import FileImageStore, ImageTooLargeError
from herpetoid.infrastructure.paths import BundlePathError, resolve_in_bundle
from herpetoid.infrastructure.pdf_export import PdfExporter


def _image(path: Path, size: tuple[int, int] = (16, 16)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    PilImage.fromarray(np.zeros((size[1], size[0], 3), np.uint8)).save(path)


# -- bundle path containment ----------------------------------------------------------------------

ESCAPING_PATHS = [
    "../secret.png",
    "images/../../secret.png",
    "..\\secret.png",
    "images\\..\\..\\secret.png",
    "/etc/passwd",
    "C:/Windows/win.ini",
    "",
    "   ",
]


@pytest.mark.parametrize("rel_path", ESCAPING_PATHS)
def test_resolve_in_bundle_rejects_escapes(tmp_path: Path, rel_path: str) -> None:
    with pytest.raises(BundlePathError):
        resolve_in_bundle(tmp_path, rel_path)


def test_resolve_in_bundle_accepts_normal_paths(tmp_path: Path) -> None:
    assert resolve_in_bundle(tmp_path, "images/abc.png") == (tmp_path / "images" / "abc.png")
    assert resolve_in_bundle(tmp_path, "thumbnails/abc.jpg").is_relative_to(tmp_path.resolve())


@pytest.mark.parametrize("rel_path", ["../secret.png", "images/../../secret.png"])
def test_image_store_load_refuses_to_read_outside_the_bundle(
    tmp_path: Path, rel_path: str
) -> None:
    """A hostile bundle must not be able to read the opener's files through an image row."""
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    _image(tmp_path / "secret.png")

    with pytest.raises(BundlePathError):
        FileImageStore(bundle).load(rel_path)


def test_image_store_save_refuses_to_write_outside_the_bundle(tmp_path: Path) -> None:
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    source = tmp_path / "source.png"
    _image(source)

    with pytest.raises(BundlePathError):
        FileImageStore(bundle).save(source, "../../pwned.png")
    assert not (tmp_path.parent / "pwned.png").exists()


def test_image_store_rejects_a_symlink_pointing_out_of_the_bundle(tmp_path: Path) -> None:
    """resolve() follows links, so a link planted inside a shared bundle is caught as an escape."""
    bundle = tmp_path / "bundle"
    (bundle / "images").mkdir(parents=True)
    outside = tmp_path / "secret.png"
    _image(outside)
    link = bundle / "images" / "link.png"
    try:
        link.symlink_to(outside)
    except (OSError, NotImplementedError):  # Windows without developer mode / privilege
        pytest.skip("symlinks not permitted in this environment")

    with pytest.raises(BundlePathError):
        FileImageStore(bundle).load("images/link.png")


def test_pdf_export_skips_thumbnails_outside_the_bundle(tmp_path: Path) -> None:
    """A dossier must never embed (and so exfiltrate) a file from outside the bundle."""
    from herpetoid.domain import Image

    bundle = tmp_path / "bundle"
    bundle.mkdir()
    _image(tmp_path / "secret.png")
    data = ExportData(
        project=Project(name="P", uuid="u", id=1),
        species=[Species(scientific_name="Calotriton asper", id=1)],
        individuals=[Individual(species_id=1, code="CA-001", id=1)],
        observations=[Observation(species_id=1, individual_id=1, id=1)],
        images=[
            Image(
                observation_id=1,
                rel_path="../secret.png",
                thumbnail_path="../secret.png",
                id=1,
            )
        ],
    )
    destination = tmp_path / "report.pdf"
    PdfExporter(bundle).export(data, destination)  # must not raise, and must not embed the file
    assert destination.exists()


# -- hostile images -------------------------------------------------------------------------------


def test_import_rejects_a_decompression_bomb(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """An image whose header declares an absurd size is refused before it is decompressed."""
    import herpetoid.infrastructure.image_store as store_module

    monkeypatch.setattr(store_module, "_MAX_PIXELS", 100)  # 10x10 budget
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    source = tmp_path / "bomb.png"
    _image(source, size=(64, 64))  # 4096 pixels > budget

    with pytest.raises(ImageTooLargeError):
        FileImageStore(bundle).import_image(source)
    assert not list(bundle.glob("images/*"))


def test_import_rejects_a_file_that_is_not_an_image(tmp_path: Path) -> None:
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    source = tmp_path / "notes.png"
    source.write_bytes(b"not an image")

    with pytest.raises(OSError):  # PIL.UnidentifiedImageError is an OSError
        FileImageStore(bundle).import_image(source)


def test_load_rejects_an_oversized_image_already_in_the_bundle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The budget also guards images arriving inside a shared bundle, which never pass import."""
    import herpetoid.infrastructure.image_store as store_module

    context = ProjectService().create(tmp_path / "bundle", "B")
    _image(tmp_path / "bundle" / "images" / "big.png", size=(64, 64))
    monkeypatch.setattr(store_module, "_MAX_PIXELS", 100)

    with pytest.raises(ImageTooLargeError):
        context.image_store.load("images/big.png")
    context.close()


# -- spreadsheet formula injection ----------------------------------------------------------------

HOSTILE = "=cmd|'/c calc'!A1"


def _hostile_data() -> ExportData:
    return ExportData(
        project=Project(name="Study", uuid="u1", id=1),
        species=[Species(scientific_name="Calotriton asper", id=1)],
        individuals=[Individual(species_id=1, code=HOSTILE, notes="+1+1", id=1)],
        observations=[
            Observation(
                species_id=1,
                individual_id=1,
                observer=HOSTILE,
                notes="@SUM(1)",
                measurements={"svl": 50.0},
                id=1,
            )
        ],
    )


def test_csv_export_neutralizes_formula_cells(tmp_path: Path) -> None:
    destination = tmp_path / "out.csv"
    CsvExporter().export(_hostile_data(), destination)
    text = destination.read_text(encoding="utf-8")

    for line in text.splitlines()[1:]:  # the header is ours, the rows are user data
        for field in line.split(","):
            assert not field.lstrip('"').startswith(("=", "+", "@")), field
    assert "'=cmd" in text  # the value is preserved, just marked as text
    assert "50.0" in text  # numbers are untouched


def test_excel_export_writes_formula_text_as_text(tmp_path: Path) -> None:
    destination = tmp_path / "out.xlsx"
    ExcelExporter().export(_hostile_data(), destination)
    workbook = load_workbook(destination)

    observer = workbook["Observations"].cell(row=2, column=4)
    assert observer.data_type == "s"  # a string, not a formula
    assert observer.value == HOSTILE  # xlsx is lossless: the value is unchanged
    code = workbook["Individuals"].cell(row=2, column=2)
    assert code.data_type == "s"
    assert code.value == HOSTILE
    assert not any(
        cell.data_type == "f" for sheet in workbook for row in sheet.iter_rows() for cell in row
    )


def test_json_export_is_left_lossless(tmp_path: Path) -> None:
    """JSON is never formula-evaluated, so it keeps the exact values for downstream analysis."""
    import json

    destination = tmp_path / "out.json"
    JsonExporter().export(_hostile_data(), destination)
    payload = json.loads(destination.read_text(encoding="utf-8"))
    assert payload["observations"][0]["observer"] == HOSTILE
    assert payload["individuals"][0]["code"] == HOSTILE
