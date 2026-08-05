"""Concrete file exporters (CSV, Excel, JSON) implementing the export port.

**Spreadsheet formula injection.** Field data is free text typed by observers, and exports are made to
be opened in Excel or LibreOffice and mailed around. A cell whose text starts with ``=``, ``+``, ``-``
or ``@`` is treated by those programs as a *formula* rather than data, which at best corrupts the
value shown to the next reader and at worst runs something on their machine. Both spreadsheet
exporters therefore neutralize such cells; the JSON exporter does not, because JSON is never
formula-evaluated and stays the lossless machine-readable path.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from herpetoid.application.export import ExportData, Exporter, observation_rows

#: Leading characters that make a spreadsheet read a cell as a formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def is_formula_like(value: object) -> bool:
    """True if a spreadsheet would interpret ``value`` as a formula rather than as text."""
    return isinstance(value, str) and value.startswith(_FORMULA_PREFIXES)


def csv_safe(value: object) -> object:
    """A CSV-safe version of ``value``: formula-like text gets a leading apostrophe.

    CSV carries no type information, so the only way to mark text as text is in the text itself.
    Numbers and dates are untouched -- this only ever fires on strings a spreadsheet would execute.
    """
    return f"'{value}" if is_formula_like(value) else value


class CsvExporter:
    format_id = "csv"

    def export(self, data: ExportData, destination: Path) -> None:
        columns, rows = observation_rows(data)
        with destination.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(
                {key: csv_safe(value) for key, value in row.items()} for row in rows
            )


class JsonExporter:
    format_id = "json"

    def export(self, data: ExportData, destination: Path) -> None:
        _columns, rows = observation_rows(data)
        payload = {
            "project": {
                "name": data.project.name,
                "uuid": data.project.uuid,
                "description": data.project.description,
            },
            "individuals": [
                {
                    "id": individual.id,
                    "code": individual.code,
                    "name": individual.name,
                    "sex": str(individual.sex),
                    "status": str(individual.status),
                    "notes": individual.notes,
                }
                for individual in data.individuals
            ],
            "observations": rows,
        }
        destination.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


class ExcelExporter:
    format_id = "xlsx"

    def export(self, data: ExportData, destination: Path) -> None:
        columns, rows = observation_rows(data)
        workbook = Workbook()
        observations_sheet = workbook.active
        observations_sheet.title = "Observations"
        observations_sheet.append(columns)
        for row in rows:
            _append_row(observations_sheet, [row.get(column) for column in columns])

        individuals_sheet = workbook.create_sheet("Individuals")
        individuals_sheet.append(["id", "code", "name", "sex", "status", "notes"])
        for individual in data.individuals:
            _append_row(
                individuals_sheet,
                [
                    individual.id,
                    individual.code,
                    individual.name,
                    str(individual.sex),
                    str(individual.status),
                    individual.notes,
                ],
            )
        workbook.save(destination)


def _append_row(sheet: Worksheet, values: list[Any]) -> None:
    """Append a row, writing formula-like text as literal text.

    Unlike CSV, xlsx cells are typed, so the value is stored **unchanged** and merely tagged as a
    string -- the export stays lossless while Excel shows the text instead of evaluating it.
    """
    sheet.append(values)
    row = sheet.max_row
    for column, value in enumerate(values, start=1):
        if is_formula_like(value):
            cell: Cell = sheet.cell(row=row, column=column)
            cell.data_type = "s"


def default_exporters() -> list[Exporter]:
    """The exporters shipped in the base app (CSV, Excel, JSON)."""
    return [CsvExporter(), ExcelExporter(), JsonExporter()]
